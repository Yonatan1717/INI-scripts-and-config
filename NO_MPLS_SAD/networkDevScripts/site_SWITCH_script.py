import os
import sys
import json
import ipaddress
import pandas as pd
from openpyxl import load_workbook

SSH_DOMAIN = "lab.local"
DEFAULT_ETHERCHANNEL_PORTS = 1
DEFAULT_ISP_VLAN = 200
# Set per variant below. Explicit Excel value always wins.
DEFAULT_SKIPPED_PORTS = 0

have_asked_rsyslog_server = False
have_asked_tacacs_server = False


def _clean_header(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def _extract_table_blocks(df):
    """Return contiguous non-empty blocks as DataFrames with first row as header."""
    blocks = []
    start = None
    non_empty = ~df.isna().all(axis=1)

    for idx, has_data in non_empty.items():
        if has_data and start is None:
            start = idx
        elif not has_data and start is not None:
            raw = df.loc[start:idx - 1].copy()
            raw = raw.dropna(axis=1, how="all")
            if not raw.empty:
                headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
                raw.columns = headers
                table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
                blocks.append((set(headers), table))
            start = None

    if start is not None:
        raw = df.loc[start:].copy().dropna(axis=1, how="all")
        if not raw.empty:
            headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
            raw.columns = headers
            table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
            blocks.append((set(headers), table))

    return blocks


def read_sheet(filename, sheet):
    """
    Reads the switch-related tables by their headers instead of fixed row numbers.
    This makes the sheet resilient to inserted/removed blank rows and extra columns.
    """
    df = pd.read_excel(filename, sheet_name=sheet, header=None)
    blocks = _extract_table_blocks(df)

    md_top = None
    md_switch = None
    swi_data = None

    for headers, table in blocks:
        if "site" in headers and "router-id" in headers:
            md_top = table
        elif "site" in headers and "secret" in headers and "router-id" not in headers:
            md_switch = table
        elif "SW" in headers:
            swi_data = table

    missing = []
    if md_top is None:
        missing.append("top metadata table (site/router-id)")
    if md_switch is None:
        missing.append("switch metadata table (site/secret)")
    if swi_data is None:
        missing.append("switch data table (SW/...)")

    if missing:
        raise ValueError(
            f"Kunne ikke finne forventede tabeller i ark '{sheet}': {', '.join(missing)}"
        )

    return {
        "md": md_switch,
        "swi_data": swi_data,
        "md_top": md_top,
    }


def _value(obj, names, default=None):
    for name in names:
        if name in obj.index:
            value = obj.get(name)
            if not pd.isna(value) and str(value).strip() != "":
                return value
    return default


def _int_value(obj, names, default=0):
    value = _value(obj, names, default)
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return int(default)


def _site_id(value):
    try:
        f = float(value)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _switch_id(value):
    return _site_id(value)


def _parse_vlan_allocation(value):
    """
    Parse format such as: 10.1-20.5-30.5-40.2
    Returns [(10,1), (20,5), ...]. VLANs with zero access ports are preserved.
    """
    if pd.isna(value) or str(value).strip() == "":
        return []

    vlan_info = []
    for part in str(value).strip().split("-"):
        part = part.strip()
        if not part:
            continue
        if "." not in part:
            raise ValueError(
                f"Ugyldig vlan-antall '{value}'. Forventet f.eks. '10.1-20.5'."
            )
        vlan_s, count_s = part.split(".", 1)
        vlan_info.append((int(float(vlan_s)), int(float(count_s))))
    return vlan_info


def _ordered_unique(values):
    seen = set()
    result = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _interface_key(prefix, ports):
    ports = list(ports)
    if not ports:
        raise ValueError("Tom portliste kan ikke gjøres om til interface-kommando")
    if len(ports) == 1:
        return f"interface {prefix}{ports[0]}"
    return f"interface range {prefix}{ports[0]} - {ports[-1]}"


def _get_isp_vlan(md):
    if not md.empty:
        value = _value(md.iloc[0], ["ISP-VLAN", "isp_vlan", "ISP VLAN"], None)
        if value is not None:
            return int(float(value))
    # Product architecture default; only matters on a switch that actually has VLAN 200.
    return DEFAULT_ISP_VLAN


def _get_etherchannel_ports(md):
    if md.empty:
        return DEFAULT_ETHERCHANNEL_PORTS
    value = _int_value(
        md.iloc[0],
        ["etherchan_num", "etherchannel_num", "etherchannel_ports", "ports_per_channel"],
        DEFAULT_ETHERCHANNEL_PORTS,
    )
    if value < 1:
        raise ValueError("etherchan_num må være minst 1")
    return value


def _is_true(value):
    """Godta True/TRUE/1/yes/ja/x fra sheets."""
    if pd.isna(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    return str(value).strip().lower() in {"true", "1", "yes", "ja", "y", "x"}


def _get_span_enabled(md):
    if md.empty:
        return False
    value = _value(md.iloc[0], ["span_port", "span port", "spanport"], None)
    return _is_true(value)


def _is_hub_site(md, md_top, site):
    """
    Determine whether this sheet represents the hub site.

    Optional Excel overrides supported:
      - is_hub / hub = TRUE/FALSE
      - role / site_role = HUB
      - hub_site = site number

    Backward-compatible product default: site 1 is the hub.
    """
    site = _site_id(site)

    for table in (md, md_top):
        if table is None or table.empty:
            continue
        row = table.iloc[0]

        explicit = _value(row, ["is_hub", "hub", "HUB"], None)
        if explicit is not None:
            return _is_true(explicit)

        role = _value(row, ["role", "site_role", "site role"], None)
        if role is not None:
            return str(role).strip().lower() == "hub"

        hub_site = _value(row, ["hub_site", "hub site"], None)
        if hub_site is not None:
            return _site_id(hub_site) == site

    return site == "1"


def _effective_vlan_info(row, md, is_hub):
    """
    VLAN 200 is the local ISP transit VLAN in this architecture.
    It is only valid on SW1 at the hub and must never be extended to
    downstream switches or spoke sites.
    """
    vlan_info = _parse_vlan_allocation(row.get("vlan-antall", ""))
    isp_vlan = _get_isp_vlan(md)
    sw_id = _switch_id(row.get("SW"))

    if not (is_hub and sw_id == "1"):
        vlan_info = [(vlan, count) for vlan, count in vlan_info if vlan != isp_vlan]

    return vlan_info


def _build_port_plan(row, md, is_hub=False):
    """
    Port model used by the workbook:

      SW1 (distribution switch):
        - skipped physical ports
        - 1 dedicated MGMT port
        - optional 1 SPAN/IDS destination port
        - access ports from low to high
        - exactly 1 physical uplink to the site router
        - num_downlink * etherchan_num ports for switch-to-switch downlinks

      SW2+ (downstream switches):
        - skipped physical ports
        - 1 dedicated MGMT port
        - access ports from low to high
        - etherchan_num ports for ONE uplink EtherChannel toward the upstream switch
        - num_downlink * etherchan_num ports for further downlinks

    This matches the Excel capacity logic:
      SW1:  ports - skipped - MGMT - SPAN? - router_uplink(1)
            - num_downlink * etherchan_num
      SW2+: ports - skipped - MGMT
            - (1 + num_downlink) * etherchan_num
    """
    has_new_num_ports = (
        "num_ports_tot" in row.index
        and not pd.isna(row.get("num_ports_tot"))
        and str(row.get("num_ports_tot")).strip() != ""
    )

    num_ports = _int_value(row, ["num_ports_tot", "num_ports"], 0)
    skipped_ports = _int_value(
        row,
        ["skiped_ports", "skipped_ports"],
        DEFAULT_SKIPPED_PORTS,
    )
    num_downlinks = _int_value(row, ["num_downlink", "num_downlinks"], 0)
    etherchan_ports = _get_etherchannel_ports(md)
    sw_id = _switch_id(row.get("SW"))
    is_primary_switch = sw_id == "1"

    if num_ports <= 0:
        raise ValueError(f"SW{sw_id}: num_ports må være > 0")
    if skipped_ports < 0:
        raise ValueError(f"SW{sw_id}: skiped_ports kan ikke være negativ")
    if num_downlinks < 0:
        raise ValueError(f"SW{sw_id}: num_downlink kan ikke være negativ")
    if has_new_num_ports and skipped_ports >= num_ports:
        raise ValueError(
            f"SW{sw_id}: skiped_ports ({skipped_ports}) må være mindre enn "
            f"num_ports_tot ({num_ports})"
        )

    first_usable = skipped_ports
    if has_new_num_ports:
        # New model: skipped ports are real physical ports that are excluded
        # from allocation and will be blackholed/shut below.
        last_port = num_ports - 1
        usable_count = num_ports - skipped_ports
        skipped_physical_ports = list(range(0, skipped_ports))
    else:
        # Legacy model: num_ports is the physical count and skipped_ports is
        # only an interface-number offset (e.g. Gi1/0/1..24).
        last_port = skipped_ports + num_ports - 1
        usable_count = num_ports
        skipped_physical_ports = []

    mgmt_port = first_usable

    # SPAN is only reserved on the primary switch in a site.
    span_enabled = is_primary_switch and _get_span_enabled(md)
    dedicated_ports = 1 + (1 if span_enabled else 0)  # MGMT (+ SPAN)

    # SW1 has one physical router uplink. Downstream switches use an
    # EtherChannel-sized uplink toward their upstream switch.
    uplink_member_count = 1 if is_primary_switch else etherchan_ports
    downlink_member_count = num_downlinks * etherchan_ports
    trunk_member_count = uplink_member_count + downlink_member_count
    available_access_slots = usable_count - dedicated_ports - trunk_member_count

    if available_access_slots < 0:
        uplink_desc = "1 router-uplink" if is_primary_switch else f"{etherchan_ports}-ports uplink"
        raise ValueError(
            f"SW{sw_id}: ikke nok porter. {usable_count} brukbare porter, men "
            f"MGMT{' + SPAN' if span_enabled else ''} + {uplink_desc} + "
            f"downlinks krever {dedicated_ports + trunk_member_count} porter."
        )

    span_port = mgmt_port + 1 if span_enabled else None
    first_access_port = mgmt_port + 1 + (1 if span_enabled else 0)
    last_access_port = first_access_port + available_access_slots - 1

    # Reserve uplink at the highest interface numbers.
    uplink_start = last_port - uplink_member_count + 1
    uplink_ports = list(range(uplink_start, last_port + 1))

    # Reserve downlink groups immediately below the uplink.
    downlink_groups = []
    first_downlink_top = uplink_start - 1
    for idx in range(num_downlinks):
        group_end = first_downlink_top - idx * etherchan_ports
        group_start = group_end - etherchan_ports + 1
        downlink_groups.append(list(range(group_start, group_end + 1)))

    # VLAN200 is only effective on hub SW1.
    vlan_info = _effective_vlan_info(row, md, is_hub)
    allocated_access_ports = sum(max(0, count) for _, count in vlan_info)
    if allocated_access_ports > available_access_slots:
        raise ValueError(
            f"SW{sw_id}: vlan-antall bruker {allocated_access_ports} access-porter, "
            f"men bare {available_access_slots} er tilgjengelige etter "
            f"MGMT/SPAN/uplink/downlinks."
        )

    expected_free = _value(row, ["num_port_ledig"], None)
    if expected_free is not None:
        try:
            expected_free = int(float(expected_free))
            if expected_free != available_access_slots:
                print(
                    f"ADVARSEL SW{sw_id}: num_port_ledig i Excel er {expected_free}, "
                    f"men generatoren beregner {available_access_slots}."
                )
        except (TypeError, ValueError):
            pass

    return {
        "num_ports": num_ports,
        "skipped_ports": skipped_ports,
        "skipped_physical_ports": skipped_physical_ports,
        "mgmt_port": mgmt_port,
        "span_port": span_port,
        "first_access_port": first_access_port,
        "last_access_port": last_access_port,
        "available_access_slots": available_access_slots,
        "allocated_access_ports": allocated_access_ports,
        "etherchan_ports": etherchan_ports,
        "is_primary_switch": is_primary_switch,
        "uplink_is_etherchannel": (not is_primary_switch and etherchan_ports > 1),
        "downlink_is_etherchannel": etherchan_ports > 1 and num_downlinks > 0,
        "num_downlinks": num_downlinks,
        "uplink_ports": uplink_ports,
        "downlink_groups": downlink_groups,
        "vlan_info": vlan_info,
    }


def create_tacacs_config(md_top):
    my_data = {"config": {}, "network_info": {}}

    global have_asked_tacacs_server, tacacs_server
    if not have_asked_tacacs_server:
        tacacs_server = input("IP-adressen til TACACS-serveren: ").strip()
        have_asked_tacacs_server = True

    tacacs_key = md_top.iloc[0].get("tacacs_key", "")

    if not tacacs_server or pd.isna(tacacs_key) or not str(tacacs_key).strip():
        raise ValueError("TACACS-server eller TACACS-key mangler")

    my_data["config"]["aaa new-model"] = []
    my_data["config"]["aaa group server tacacs+ TACACS-GROUP"] = [
        f"server-private {tacacs_server} key {str(tacacs_key).strip()}",
        "ip tacacs source-interface Vlan10",
        "exit",
    ]
    my_data["config"]["aaa authentication login default group TACACS-GROUP local"] = []
    my_data["config"]["aaa authorization exec default group TACACS-GROUP local"] = []

    return my_data


def create_rsyslog_config():
    my_data = {"config": {}, "network_info": {}}

    global have_asked_rsyslog_server, rsyslog_server
    if not have_asked_rsyslog_server:
        rsyslog_server = input("IP-adressen til Rsyslog-serveren: ").strip()
        have_asked_rsyslog_server = True

    if not rsyslog_server:
        raise ValueError("Rsyslog-server mangler")

    my_data["config"]["service timestamps log datetime msec show-timezone"] = []
    my_data["config"][f"logging host {rsyslog_server} transport udp port 514"] = []
    my_data["config"]["logging trap informational"] = []
    my_data["config"]["logging source-interface Vlan10"] = []

    return my_data


def enable_ssh(md, domain=SSH_DOMAIN, mgmt_network=None, mgmt_wildcard=None):
    my_data = {"config": {}, "network_info": {}}
    if md.empty:
        return my_data

    row = md.iloc[0]
    username = row.get("brukernavn", "")
    password = row.get("passord", "")
    vty_lines = row.get("vty_lines", "0-4")

    if pd.isna(username) or pd.isna(password):
        return my_data

    username = str(username).strip()
    password = str(password).strip()
    if not username or not password:
        return my_data

    my_data["config"][f"ip domain name {domain}"] = []
    my_data["config"][f"username {username} privilege 15 secret 9 {password}"] = []
    my_data["config"]["crypto key generate rsa general-keys modulus 4096"] = []
    my_data["config"]["ip ssh version 2"] = []

    vty_cfg = ["login authentication default", "transport input ssh"]
    if mgmt_network and mgmt_wildcard:
        my_data["config"]["ip access-list standard SSH-MGMT-ONLY"] = [
            f"permit {mgmt_network} {mgmt_wildcard}",
            "exit",
        ]
        vty_cfg.insert(0, "access-class SSH-MGMT-ONLY in")
    vty_cfg.append("exit")

    my_data["config"][f"line vty {' '.join(x.strip() for x in str(vty_lines).split('-'))}"] = vty_cfg

    return my_data


def global_config(md, md_top, swi_data, is_hub):
    info = {"config": {}, "network_info": {}}
    site = _site_id(md.iloc[0]["site"])
    secret = md.iloc[0].get("secret", "")

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        mgmt_ip = row["MGMT ip"]
        gateway = row["gateway"]
        mask = row["mask"]
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, is_hub)

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg[f"hostname {sw_name}"] = []
        sw_cfg[f"enable secret 9 {secret}"] = []
        sw_cfg.update(create_tacacs_config(md_top)["config"])
        mgmt_network = mgmt_wildcard = None
        try:
            mgmt_net = ipaddress.ip_network(f"{mgmt_ip}/{mask}", strict=False)
            mgmt_network, mgmt_wildcard = str(mgmt_net.network_address), str(mgmt_net.hostmask)
        except (ValueError, TypeError):
            pass
        sw_cfg.update(enable_ssh(md, mgmt_network=mgmt_network, mgmt_wildcard=mgmt_wildcard)["config"])
        sw_cfg.update(create_rsyslog_config()["config"])

        sw_cfg[f"vlan {mgmt_vlan}"] = [f"name MGMT_VLAN_{mgmt_vlan}", "exit"]
        sw_cfg[f"interface vlan {mgmt_vlan}"] = [
            f"ip address {mgmt_ip} {mask}",
            "no shutdown",
            "exit",
        ]

        sw_cfg[f"interface {intf_prefix}{plan['mgmt_port']}"] = [
            f"description Dedicated management access port for VLAN {mgmt_vlan}",
            # Trusted: physically controlled infra port for statically addressed mgmt hosts
            # (e.g. the TACACS/Rsyslog server), which have no DHCP snooping binding for DAI to check.
            "ip arp inspection trust",
            "switchport mode access",
            f"switchport access vlan {mgmt_vlan}",
            "switchport port-security",
            "switchport port-security maximum 2",
            "switchport port-security violation restrict",
            "spanning-tree bpduguard enable",
            "spanning-tree portfast",
            "no shutdown",
            "exit",
        ]

        if plan["span_port"] is not None:
            span_vlans = _ordered_unique([mgmt_vlan, *[v for v, _ in plan["vlan_info"]]])
            sw_cfg[f"interface {intf_prefix}{plan['span_port']}"] = [
                "description Dedicated SPAN destination port for IDS/IPS",
                "no shutdown",
                "exit",
            ]
            sw_cfg[f"monitor session 1 source vlan {','.join(map(str, span_vlans))} both"] = []
            sw_cfg[f"monitor session 1 destination interface {intf_prefix}{plan['span_port']}"] = []

        sw_cfg[f"ip default-gateway {gateway}"] = []
        sw_cfg[f"ntp server {gateway}"] = []

    return info


def config_vlan(swi_data, site, md, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, is_hub)
        vlan_info = plan["vlan_info"]

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg["vlan 999"] = ["name NATIVE_UBRUKT", "exit"]

        current_port = plan["first_access_port"]
        for vlan, count in vlan_info:
            # VLAN must exist even when it has zero local access ports.
            sw_cfg[f"vlan {vlan}"] = [f"name VLAN_{vlan}", "exit"]

            if count <= 0:
                continue

            ports = list(range(current_port, current_port + count))
            current_port += count
            key = _interface_key(intf_prefix, ports)

            port_cfg = [
                f"description access port for VLAN {vlan}",
            ]
            if vlan == isp_vlan:
                # ISP transit port: infrastructure hand-off, not a client access port.
                # Excluded from DAI/DHCP snooping VLAN checks, so trust is meaningless here;
                # port-security/bpduguard are client-host protections that don't apply to it either.
                port_cfg.extend([
                    "switchport mode access",
                    f"switchport access vlan {vlan}",
                    "no shutdown",
                    "exit",
                ])
                sw_cfg[key] = port_cfg
                continue

            port_cfg.extend([
                "switchport mode access",
                f"switchport access vlan {vlan}",
                "switchport port-security",
                "switchport port-security maximum 2",
                "switchport port-security violation restrict",
                "spanning-tree bpduguard enable",
                "spanning-tree portfast",
                "no shutdown",
                "exit",
            ])
            sw_cfg[key] = port_cfg

    return info


def _trunk_config(vlans, description, channel_group=None, trusted=True):
    lines = [
        description,
        "switchport trunk encapsulation dot1q",
        "switchport trunk native vlan 999",
        "switchport mode trunk",
        f"switchport trunk allowed vlan {','.join(map(str, vlans))}",
    ]
    if trusted:
        # Trust only flows toward the DHCP server / infrastructure side (uplinks).
        lines.extend(["ip dhcp snooping trust", "ip arp inspection trust"])
    if channel_group is not None:
        lines.append(f"channel-group {channel_group} mode active")
    lines.extend(["no shutdown", "exit"])
    return lines


def config_trunk_and_dchp_snooping(swi_data, site, md, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, is_hub)

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        vlan_numbers = [vlan for vlan, _ in plan["vlan_info"]]
        all_trunk_vlans = _ordered_unique([mgmt_vlan, *vlan_numbers, 999])

        # VLAN 200 is local ISP transit in this architecture and must not be extended downstream.
        downlink_vlans = [v for v in all_trunk_vlans if v != isp_vlan]

        # No snooping/DAI on ISP transit or blackhole/native VLAN 999.
        inspection_vlans = [v for v in all_trunk_vlans if v not in {isp_vlan, 999}]
        if inspection_vlans:
            sw_cfg["ip dhcp snooping"] = []
            sw_cfg[f"ip dhcp snooping vlan {','.join(map(str, inspection_vlans))}"] = []
            sw_cfg["no ip dhcp snooping information option"] = []
            sw_cfg[f"ip arp inspection vlan {','.join(map(str, inspection_vlans))}"] = []

        # UPLINK:
        # - SW1: exactly one physical trunk to the site router.
        # - SW2+: EtherChannel-sized uplink toward the upstream switch when
        #   etherchan_num > 1; otherwise a normal physical trunk.
        if plan["uplink_is_etherchannel"]:
            uplink_po = 1
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                downlink_vlans,
                f"description UPLINK EtherChannel member(s) - Port-channel{uplink_po}",
                channel_group=uplink_po,
                trusted=True,
            )
            sw_cfg[f"interface Port-channel{uplink_po}"] = _trunk_config(
                downlink_vlans,
                f"description UPLINK Port-channel{uplink_po} toward upstream switch "
                f"for VLAN {','.join(map(str, downlink_vlans))}",
                trusted=True,
            )
            first_downlink_channel = 2
        else:
            uplink_vlans = all_trunk_vlans if plan["is_primary_switch"] else downlink_vlans
            uplink_desc = (
                "UPLINK trunk to site router"
                if plan["is_primary_switch"]
                else "UPLINK trunk toward upstream switch"
            )
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                uplink_vlans,
                f"description {uplink_desc} for VLAN {','.join(map(str, uplink_vlans))}",
                trusted=True,
            )
            first_downlink_channel = 1

        # DOWNLINKS:
        # Channel-group numbers are local. SW1 starts with Po1. On downstream
        # switches Po1 is reserved for the uplink, so their downlinks start at Po2.
        if plan["downlink_is_etherchannel"]:
            for offset, ports in enumerate(plan["downlink_groups"]):
                channel_id = first_downlink_channel + offset
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    downlink_vlans,
                    f"description DOWNLINK EtherChannel member(s) - Port-channel{channel_id}",
                    channel_group=channel_id,
                    trusted=False,
                )
                sw_cfg[f"interface Port-channel{channel_id}"] = _trunk_config(
                    downlink_vlans,
                    f"description DOWNLINK Port-channel{channel_id} for VLAN "
                    f"{','.join(map(str, downlink_vlans))}",
                    trusted=False,
                )
        else:
            for idx, ports in enumerate(plan["downlink_groups"], start=1):
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    downlink_vlans,
                    f"description DOWNLINK trunk {idx} for VLAN "
                    f"{','.join(map(str, downlink_vlans))}",
                    trusted=False,
                )

        # New num_ports_tot model: skipped ports are real physical interfaces,
        # so explicitly blackhole/shut them instead of leaving them in VLAN 1.
        if plan["skipped_physical_ports"]:
            sw_cfg[_interface_key(intf_prefix, plan["skipped_physical_ports"])] = [
                "description UBRUKT - SKIPPED/RESERVED",
                "switchport mode access",
                "switchport access vlan 999",
                "shutdown",
                "exit",
            ]

        # Remaining unassigned access slots are blackholed and shut down.
        unused_start = plan["first_access_port"] + plan["allocated_access_ports"]
        unused_end = plan["last_access_port"]
        if unused_start <= unused_end:
            unused_ports = list(range(unused_start, unused_end + 1))
            sw_cfg[_interface_key(intf_prefix, unused_ports)] = [
                "description UBRUKT - BLACKHOLE VLAN",
                "switchport mode access",
                "switchport access vlan 999",
                "shutdown",
                "exit",
            ]

    return info


def update_site_config(data, swi_data, sn, conf):
    for raw_sw_id in swi_data["SW"].dropna().unique():
        sw_id = _switch_id(raw_sw_id)
        sw_name = f"SW{sw_id}-SITE-{sn}"
        data[f"site {sn}"]["config"].setdefault(sw_name, {})
        data[f"site {sn}"]["config"][sw_name].update(conf["config"][sw_name])
    return data


def fetch_site_data(config_file):
    try:
        with open(config_file, "r") as f:
            try:
                return json.load(f)
            except json.JSONDecodeError:
                return {}
    except FileNotFoundError:
        return {}


def _refresh_ssh_acls(data, swi_data, sn):
    """
    Record this site's MGMT subnet and re-sync the SSH-MGMT-ONLY ACL on every
    switch across every site processed so far, so newly added sites are
    automatically permitted everywhere (not just locally).
    """
    mgmt_networks = data.setdefault("_mgmt_networks", {})

    row = swi_data.iloc[0]
    try:
        net = ipaddress.ip_network(f"{row.get('MGMT ip')}/{row.get('mask')}", strict=False)
        mgmt_networks[sn] = [str(net.network_address), str(net.hostmask)]
    except (ValueError, TypeError):
        pass

    if not mgmt_networks:
        return

    acl_lines = [f"permit {network} {wildcard}" for network, wildcard in sorted(mgmt_networks.values())]
    acl_lines.append("exit")

    for site_key, site_val in data.items():
        if site_key == "_mgmt_networks":
            continue
        for sw_cfg in site_val.get("config", {}).values():
            if "ip access-list standard SSH-MGMT-ONLY" in sw_cfg:
                sw_cfg["ip access-list standard SSH-MGMT-ONLY"] = acl_lines


def create_site_sw_config(file, sheet, config_file):
    sheet_data = read_sheet(file, sheet)
    data = fetch_site_data(config_file)

    md = sheet_data["md"]
    md_top = sheet_data["md_top"]
    swi_data = sheet_data["swi_data"]

    sn = _site_id(md.iloc[0]["site"])
    is_hub = _is_hub_site(md, md_top, sn)
    data[f"site {sn}"] = {"config": {}}

    vlan_conf = config_vlan(swi_data, sn, md, is_hub)
    data = update_site_config(data, swi_data, sn, vlan_conf)

    global_conf = global_config(md, md_top, swi_data, is_hub)
    data = update_site_config(data, swi_data, sn, global_conf)

    trunk_conf = config_trunk_and_dchp_snooping(swi_data, sn, md, is_hub)
    data = update_site_config(data, swi_data, sn, trunk_conf)

    _refresh_ssh_acls(data, swi_data, sn)

    with open(config_file, "w") as f:
        json.dump(data, f, indent=4)

    return data


def config_to_text(data, indent=0):
    lines = []
    prefix = "    " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            lines.append("!")
            lines.append(prefix + key)
            lines.extend(config_to_text(value, indent + 1))
    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                lines.append(prefix + value)
            else:
                lines.extend(config_to_text(value, indent))
    elif isinstance(data, str):
        lines.append(prefix + data)

    return lines


def create_or_update_config_files(data):
    os.makedirs("siteSwichTextConfigs", exist_ok=True)

    for site, site_data in data.items():
        if site == "_mgmt_networks":
            continue
        site_dir = f"siteSwichTextConfigs/{site}"
        os.makedirs(site_dir, exist_ok=True)

        for sw_name, config in site_data["config"].items():
            text = config_to_text(config)
            with open(f"{site_dir}/{sw_name}.txt", "w", encoding="utf-8") as f:
                f.write("\n".join(text))

    print("Text versjon av switch-configene er lagret i siteSwichTextConfigs/")


def create_sw_configs_main(file, config_file="site_switch_config.json"):
    sites_sheets = load_workbook(file, read_only=True).sheetnames

    data = {}
    for sheet in sites_sheets:
        data = create_site_sw_config(file, sheet, config_file)

    create_or_update_config_files(data)


def main():
    file = sys.argv[1]
    config_file = "site_switch_config.json" if len(sys.argv) < 3 else sys.argv[2]
    create_sw_configs_main(file, config_file)


if __name__ == "__main__":
    main()
