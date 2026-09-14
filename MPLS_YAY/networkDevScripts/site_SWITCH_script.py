import os
import sys
import json
import ipaddress
import pandas as pd
from openpyxl import load_workbook

SSH_DOMAIN = "lab.local"
DEFAULT_ETHERCHANNEL_PORTS = 1
DEFAULT_ISP_VLAN = 200
# Set per variant below. Explicit Excel value always wins.
DEFAULT_SKIPPED_PORTS = 0
DEFAULT_RSPAN_VLAN = 900
VALID_SPAN_MODES = {"SPAN", "RSPAN", "ERSPAN"}
MONITORING_VRF = "MONITORING"
DEFAULT_ERSPAN_ID = 100
REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN = False


def _clean_header(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def _extract_table_blocks(df):
    """Return contiguous non-empty blocks as DataFrames with first row as header."""
    blocks = []
    start = None
    non_empty = ~df.isna().all(axis=1)

    for idx, has_data in non_empty.items():
        if has_data and start is None:
            start = idx
        elif not has_data and start is not None:
            raw = df.loc[start:idx - 1].copy()
            raw = raw.dropna(axis=1, how="all")
            if not raw.empty:
                headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
                raw.columns = headers
                table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
                blocks.append((set(headers), table))
            start = None

    if start is not None:
        raw = df.loc[start:].copy().dropna(axis=1, how="all")
        if not raw.empty:
            headers = [_clean_header(x) for x in raw.iloc[0].tolist()]
            raw.columns = headers
            table = raw.iloc[1:].dropna(how="all").reset_index(drop=True)
            blocks.append((set(headers), table))

    return blocks


def read_sheet(filename, sheet):
    """
    Reads the switch-related tables by their headers instead of fixed row numbers.
    This makes the sheet resilient to inserted/removed blank rows and extra columns.
    """
    df = pd.read_excel(filename, sheet_name=sheet, header=None)
    blocks = _extract_table_blocks(df)

    md_top = None
    md_switch = None
    swi_data = None
    ip_data = None
    vrf_data = None
    tunnel_data = None

    for headers, table in blocks:
        if "site" in headers and "router-id" in headers:
            md_top = table
        elif "site" in headers and "secret" in headers and "router-id" not in headers:
            md_switch = table
        elif "SW" in headers:
            swi_data = table
        elif {"vrf", "vlan", "mask", "address min"}.issubset(headers):
            ip_data = table
        elif {"vrf", "loopback", "laddr"}.issubset(headers):
            vrf_data = table
        elif {"tunnel id", "vrf"}.issubset(headers):
            tunnel_data = table

    missing = []
    if md_top is None:
        missing.append("top metadata table (site/router-id)")
    if md_switch is None:
        missing.append("switch metadata table (site/secret)")
    if swi_data is None:
        missing.append("switch data table (SW/...)")

    if missing:
        raise ValueError(
            f"Kunne ikke finne forventede tabeller i ark '{sheet}': {', '.join(missing)}"
        )

    return {
        "md": md_switch,
        "swi_data": swi_data,
        "md_top": md_top,
        "ip_data": ip_data,
        "vrf_data": vrf_data,
        "tunnel_data": tunnel_data,
    }


def _value(obj, names, default=None):
    for name in names:
        if name in obj.index:
            value = obj.get(name)
            if not pd.isna(value) and str(value).strip() != "":
                return value
    return default


def _int_value(obj, names, default=0):
    value = _value(obj, names, default)
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return int(default)


def _site_id(value):
    try:
        f = float(value)
        if f.is_integer():
            return str(int(f))
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _switch_id(value):
    return _site_id(value)


def _parse_vlan_allocation(value):
    """
    Parse format such as: 10.1-20.5-30.5-40.2
    Returns [(10,1), (20,5), ...]. VLANs with zero access ports are preserved.
    """
    if pd.isna(value) or str(value).strip() == "":
        return []

    vlan_info = []
    for part in str(value).strip().split("-"):
        part = part.strip()
        if not part:
            continue
        if "." not in part:
            raise ValueError(
                f"Ugyldig vlan-antall '{value}'. Forventet f.eks. '10.1-20.5'."
            )
        vlan_s, count_s = part.split(".", 1)
        vlan_info.append((int(float(vlan_s)), int(float(count_s))))
    return vlan_info


def _ordered_unique(values):
    seen = set()
    result = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _interface_key(prefix, ports):
    ports = list(ports)
    if not ports:
        raise ValueError("Tom portliste kan ikke gjøres om til interface-kommando")
    if len(ports) == 1:
        return f"interface {prefix}{ports[0]}"
    return f"interface range {prefix}{ports[0]} - {ports[-1]}"


def _source_interface_ranges(ports, intf_prefix, direction="rx"):
    """Compress physical interface numbers into Cisco monitor-session ranges.

    Examples:
      [1]       -> source interface g0/1 rx
      [1,2,3]   -> source interface g0/1 - 3 rx
      [1,2,5,6] -> source interface g0/1 - 2 rx
                   source interface g0/5 - 6 rx
    """
    ports = sorted(set(int(p) for p in ports))
    if not ports:
        return []

    lines = []
    run_start = ports[0]
    run_end = ports[0]

    def emit(start, end):
        if start == end:
            return f"source interface {intf_prefix}{start} {direction}"
        return f"source interface {intf_prefix}{start} - {end} {direction}"

    for port in ports[1:]:
        if port == run_end + 1:
            run_end = port
            continue

        lines.append(emit(run_start, run_end))
        run_start = run_end = port

    lines.append(emit(run_start, run_end))
    return lines


def _erspan_source_interface_lines(plan, intf_prefix):
    """Build low-duplication ERSPAN sources for one switch.

    - Client traffic is mirrored only when it ENTERS a client access port (rx).
      The same frame is therefore not mirrored again while traversing
      switch-to-switch trunks.
    - Contiguous client access ports are emitted as a Cisco source-interface
      range to keep the generated configuration compact.
    - On SW1 at a spoke site, the router-facing uplink is also mirrored rx.
      This captures traffic entering the LAN from the router/WAN, so Suricata
      still sees both directions of routed client flows.
    - MGMT/server ports, monitoring/sensor ports, downlinks, unused ports and
      downstream-switch uplinks are deliberately excluded.
    """
    lines = []

    # Access ports are allocated contiguously in the current port model, but
    # use the generic range helper so the output also stays correct if gaps are
    # introduced later.
    access_start = plan["first_access_port"]
    access_ports = list(
        range(access_start, access_start + plan["allocated_access_ports"])
    )
    lines.extend(_source_interface_ranges(access_ports, intf_prefix, "rx"))

    # SW1 has exactly one physical router-facing uplink in this architecture.
    # Keep it separate from the access-port range even if interface numbers
    # happen to be adjacent: it has a different monitoring purpose.
    if plan["is_primary_switch"]:
        lines.extend(_source_interface_ranges(plan["uplink_ports"], intf_prefix, "rx"))

    return lines


def _get_isp_vlan(md):
    if not md.empty:
        value = _value(md.iloc[0], ["ISP-VLAN", "isp_vlan", "ISP VLAN"], None)
        if value is not None:
            return int(float(value))
    # Product architecture default; only matters on a switch that actually has VLAN 200.
    return DEFAULT_ISP_VLAN


def _get_etherchannel_ports(md):
    if md.empty:
        return DEFAULT_ETHERCHANNEL_PORTS
    value = _int_value(
        md.iloc[0],
        ["etherchan_num", "etherchannel_num", "etherchannel_ports", "ports_per_channel"],
        DEFAULT_ETHERCHANNEL_PORTS,
    )
    if value < 1:
        raise ValueError("etherchan_num må være minst 1")
    return value


def _is_true(value):
    """Godta True/TRUE/1/yes/ja/x fra sheets."""
    if pd.isna(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    return str(value).strip().lower() in {"true", "1", "yes", "ja", "y", "x"}


def _get_span_mode(md):
    """
    Returner valgt speilingsmodus for sitet.

    Gyldige verdier i Excel:
      - SPAN
      - RSPAN
      - ERSPAN
      - tom / NONE / OFF / INGEN = deaktivert
    """
    if md.empty:
        return None

    value = _value(md.iloc[0], ["span_mode", "span mode", "spanmode"], None)
    if value is None:
        return None

    mode = str(value).strip().upper()
    if mode in {"", "NONE", "OFF", "INGEN", "DISABLED", "FALSE", "0"}:
        return None

    if mode not in VALID_SPAN_MODES:
        raise ValueError(
            f"Ugyldig span_mode '{value}'. Gyldige verdier er "
            f"{', '.join(sorted(VALID_SPAN_MODES))}, eller tom celle for ingen speiling."
        )
    return mode


def _get_rspan_vlan(md):
    if md.empty:
        return DEFAULT_RSPAN_VLAN
    vlan = _int_value(md.iloc[0], ["rspan_vlan", "rspan vlan"], DEFAULT_RSPAN_VLAN)
    if not 1 <= vlan <= 4094 or vlan == 999:
        raise ValueError(f"Ugyldig RSPAN-VLAN {vlan}. Velg VLAN 1-4094, men ikke 999.")
    return vlan


def _get_management_server_ip(md_top, names, label):
    """Read and validate a management-service IP from the top Excel metadata."""
    if md_top is None or md_top.empty:
        raise ValueError(f"{label}-server mangler i Excel-metadata.")

    value = _value(md_top.iloc[0], names, None)
    if value is None:
        raise ValueError(
            f"{label}-server mangler i Excel. Forventet felt: {names[0]}."
        )

    try:
        return str(ipaddress.ip_address(str(value).strip()))
    except ValueError as exc:
        raise ValueError(f"Ugyldig {label}-server-IP i Excel: {value}") from exc


def _get_tacacs_server_ip(md_top):
    return _get_management_server_ip(
        md_top,
        ["tacacs_server_ip", "tacacs server ip", "tacacs_server"],
        "TACACS",
    )


def _get_syslog_server_ip(md_top):
    return _get_management_server_ip(
        md_top,
        ["syslog_server_ip", "rsyslog_server_ip", "syslog server ip", "syslog_server"],
        "Syslog",
    )


def _find_vrf_rows(df, vrf_name):
    if df is None or df.empty or "vrf" not in df.columns:
        return df.iloc[0:0] if df is not None else pd.DataFrame()
    names = df["vrf"].astype(str).str.strip().str.upper()
    return df[names == vrf_name.upper()]


def _get_monitoring_service(ip_data, vrf_data, site):
    """Return the MONITORING service definition used by ERSPAN.

    NO-MPLS requires an explicit MONITORING VRF/IP row because it also needs
    Tunnel50/DMVPN transport.  MPLS can derive the fixed reference-architecture
    MONITORING network automatically when those rows are absent:
      Site N -> VLAN 50, 10.(50+N).0.0/24, gateway .1.
    """
    vrf_rows = _find_vrf_rows(vrf_data, MONITORING_VRF)
    ip_rows = _find_vrf_rows(ip_data, MONITORING_VRF)

    if vrf_rows.empty or ip_rows.empty:
        if not REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN:
            try:
                site_no = int(float(site))
                second_octet = 50 + site_no
                if not 0 <= second_octet <= 255:
                    raise ValueError
                network = ipaddress.ip_network(f"10.{second_octet}.0.0/24")
                return {
                    "vrf": MONITORING_VRF,
                    "vlan": 50,
                    "mask": "255.255.255.0",
                    "gateway": str(network.network_address + 1),
                    "network": network,
                }
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Site {site}: kunne ikke utlede MONITORING-nett automatisk for MPLS."
                ) from exc

        if vrf_rows.empty:
            raise ValueError(
                f"Site {site}: span_mode=ERSPAN krever VRF '{MONITORING_VRF}' "
                "i Excel sin VRF-tabell for NO-MPLS router-transport."
            )
        raise ValueError(
            f"Site {site}: span_mode=ERSPAN krever en egen VLAN/subnett-rad med "
            f"vrf={MONITORING_VRF} i IP-tabellen."
        )

    if len(vrf_rows) != 1 or len(ip_rows) != 1:
        raise ValueError(
            f"Site {site}: ERSPAN forventer nøyaktig en {MONITORING_VRF}-rad i "
            "VRF-tabellen og en i IP/VLAN-tabellen."
        )

    row = ip_rows.iloc[0]
    try:
        vlan = int(float(row["vlan"]))
        mask = str(row["mask"]).strip()
        gateway = str(row.get("gateway", row.get("address min", ""))).strip()
        network_value = row.get("nett id", gateway)
        network = ipaddress.ip_network(f"{network_value}/{mask}", strict=False)
        gateway_ip = ipaddress.ip_address(gateway)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Site {site}: ugyldig MONITORING VLAN/subnett/gateway i Excel.") from exc

    if not 1 <= vlan <= 4094 or vlan == 999:
        raise ValueError(f"Site {site}: ugyldig MONITORING VLAN {vlan}.")
    if gateway_ip not in network:
        raise ValueError(
            f"Site {site}: MONITORING gateway {gateway} ligger ikke i {network}."
        )

    return {
        "vrf": MONITORING_VRF,
        "vlan": vlan,
        "mask": mask,
        "gateway": gateway,
        "network": network,
    }


def _get_switch_monitoring_ip(row, monitoring, site, sw_id):
    value = _value(
        row,
        ["MONITORING ip", "monitoring ip", "MONITORING IP", "monitoring_ip"],
        None,
    )
    if value is None:
        raise ValueError(
            f"Site {site} SW{sw_id}: span_mode=ERSPAN krever en egen 'MONITORING ip' "
            "i switch-tabellen. MGMT ip kan ikke brukes som ERSPAN origin."
        )

    try:
        address = ipaddress.ip_address(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"Site {site} SW{sw_id}: ugyldig MONITORING ip: {value}") from exc

    network = monitoring["network"]
    if address not in network or address in {network.network_address, network.broadcast_address}:
        raise ValueError(
            f"Site {site} SW{sw_id}: MONITORING ip {address} er ikke en gyldig host i {network}."
        )
    if str(address) == monitoring["gateway"]:
        raise ValueError(
            f"Site {site} SW{sw_id}: MONITORING ip kan ikke være gateway {monitoring['gateway']}."
        )
    return str(address)


def _validate_erspan_architecture_for_workbook(file, sheets):
    """Validate ERSPAN and derive the destination from HUB SW1 MONITORING IP.

    The user does not enter an ERSPAN destination manually.  In ERSPAN mode,
    HUB SW1 terminates the GRE/ERSPAN stream and forwards the decapsulated
    mirrored frames out its dedicated physical monitor destination port.
    """
    records = []
    hubs = []
    active_erspan_sites = set()

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        md_top = sheet_data["md_top"]
        site = _site_id(md.iloc[0]["site"])
        ip_data = sheet_data.get("ip_data")
        vrf_data = sheet_data.get("vrf_data")
        tunnel_data = sheet_data.get("tunnel_data")
        mode = _get_span_mode(md)
        is_hub = _is_hub_site(md, md_top, site)

        record = {
            "sheet": sheet,
            "site": site,
            "data": sheet_data,
            "mode": mode,
            "is_hub": is_hub,
        }
        records.append(record)
        if is_hub:
            hubs.append(record)
        if mode == "ERSPAN":
            active_erspan_sites.add(site)

    if not active_erspan_sites:
        return {
            "active_sites": set(),
            "hub_site": None,
            "destination_ip": None,
            "erspan_id": DEFAULT_ERSPAN_ID,
        }

    if len(hubs) != 1:
        raise ValueError(
            f"ERSPAN krever nøyaktig ett HUB-site. Fant {len(hubs)} HUB-sites."
        )

    hub = hubs[0]
    if hub["mode"] != "ERSPAN":
        raise ValueError(
            f"ERSPAN er aktivert på site(s) {sorted(active_erspan_sites)}, men HUB Site "
            f"{hub['site']} har span_mode={hub['mode'] or 'OFF'}. HUB må også stå i ERSPAN."
        )

    hub_data = hub["data"]
    hub_monitoring = _get_monitoring_service(
        hub_data.get("ip_data"), hub_data.get("vrf_data"), hub["site"]
    )
    hub_sw1_rows = hub_data["swi_data"][
        hub_data["swi_data"]["SW"].apply(_switch_id) == "1"
    ]
    if len(hub_sw1_rows) != 1:
        raise ValueError(
            f"HUB Site {hub['site']}: ERSPAN krever nøyaktig én SW1 i switch-tabellen."
        )
    destination = _get_switch_monitoring_ip(
        hub_sw1_rows.iloc[0], hub_monitoring, hub["site"], "1"
    )

    # Destination must be the HUB SW1 MONITORING address and never a MGMT address.
    mgmt_rows = _find_vrf_rows(hub_data.get("ip_data"), "MGMT")
    if not mgmt_rows.empty:
        r = mgmt_rows.iloc[0]
        mgmt_net = ipaddress.ip_network(
            f"{r.get('nett id', r.get('address min'))}/{r['mask']}", strict=False
        )
        if ipaddress.ip_address(destination) in mgmt_net:
            raise ValueError(
                f"HUB SW1 ERSPAN-destinasjon {destination} ligger i MGMT-nettet {mgmt_net}."
            )

    # Validate every active site's MONITORING transport and unique switch origin IPs.
    globally_seen_monitoring_ips = set()
    for record in records:
        if record["mode"] != "ERSPAN":
            continue

        site = record["site"]
        sheet_data = record["data"]
        monitoring = _get_monitoring_service(
            sheet_data.get("ip_data"), sheet_data.get("vrf_data"), site
        )

        for _, sw_row in sheet_data["swi_data"].iterrows():
            sw_id = _switch_id(sw_row["SW"])
            mon_ip = _get_switch_monitoring_ip(sw_row, monitoring, site, sw_id)
            if mon_ip in globally_seen_monitoring_ips:
                raise ValueError(
                    f"MONITORING ip {mon_ip} er brukt av flere switcher i arbeidsboken."
                )
            globally_seen_monitoring_ips.add(mon_ip)

        if REQUIRE_MONITORING_TUNNEL_FOR_ERSPAN:
            tun_rows = _find_vrf_rows(sheet_data.get("tunnel_data"), MONITORING_VRF)
            if tun_rows.empty:
                raise ValueError(
                    f"Site {site}: NO-MPLS + ERSPAN krever en egen DMVPN/EIGRP-tunnel "
                    f"for VRF {MONITORING_VRF}. Legg til en MONITORING-rad i tunnel-tabellen."
                )

    return {
        "active_sites": active_erspan_sites,
        "hub_site": hub["site"],
        "destination_ip": destination,
        "erspan_id": DEFAULT_ERSPAN_ID,
    }


def _validate_management_servers_for_workbook(file, sheets):
    """Ensure TACACS/Syslog server IPs are explicit, consistent and in HUB MGMT."""
    values = []
    hub_record = None

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        md_top = sheet_data["md_top"]
        site = _site_id(md.iloc[0]["site"])
        tacacs = _get_tacacs_server_ip(md_top)
        syslog = _get_syslog_server_ip(md_top)
        values.append((site, tacacs, syslog))
        if _is_hub_site(md, md_top, site):
            if hub_record is not None:
                raise ValueError("Flere HUB-sites funnet ved validering av management-servere.")
            hub_record = (site, sheet_data, tacacs, syslog)

    if hub_record is None:
        raise ValueError("Fant ikke HUB-site ved validering av TACACS/Syslog-servere.")

    tacacs_values = {x[1] for x in values}
    syslog_values = {x[2] for x in values}
    if len(tacacs_values) != 1 or len(syslog_values) != 1:
        detail = ", ".join(
            f"Site {site}: TACACS={tacacs}, Syslog={syslog}"
            for site, tacacs, syslog in values
        )
        raise ValueError(
            "TACACS/Syslog-serverne må være konsistente mellom site-arkene. " + detail
        )

    site, sheet_data, tacacs, syslog = hub_record
    mgmt_rows = _find_vrf_rows(sheet_data.get("ip_data"), "MGMT")
    if mgmt_rows.empty:
        raise ValueError(f"HUB Site {site}: finner ikke MGMT-subnett for servervalidering.")
    r = mgmt_rows.iloc[0]
    mgmt_net = ipaddress.ip_network(
        f"{r.get('nett id', r.get('address min'))}/{r['mask']}", strict=False
    )
    for label, server in (("TACACS", tacacs), ("Syslog", syslog)):
        addr = ipaddress.ip_address(server)
        if addr not in mgmt_net or addr in {mgmt_net.network_address, mgmt_net.broadcast_address}:
            raise ValueError(
                f"{label}-server {server} må ligge som gyldig host i HUB MGMT-nettet {mgmt_net}."
            )

    return {"tacacs": tacacs, "syslog": syslog, "same_server": tacacs == syslog}


def _validate_span_modes_for_workbook(file, sheets):
    """
    Alle sites som har speiling aktivert må bruke samme modus.
    Tom span_mode betyr at speiling er deaktivert på det sitet og regnes
    derfor ikke som en konflikt.
    """
    modes_by_site = {}
    enabled_modes = set()

    for sheet in sheets:
        sheet_data = read_sheet(file, sheet)
        md = sheet_data["md"]
        site = _site_id(md.iloc[0]["site"])
        mode = _get_span_mode(md)
        modes_by_site[site] = mode
        if mode is not None:
            enabled_modes.add(mode)

    if len(enabled_modes) > 1:
        details = ", ".join(
            f"Site {site}={mode or 'OFF'}"
            for site, mode in sorted(modes_by_site.items())
        )
        raise ValueError(
            "Konflikt i span_mode mellom sites. Aktive sites må bruke samme "
            f"speilingsmodus. Fant: {details}"
        )

    return modes_by_site


def _is_hub_site(md, md_top, site):
    """
    Determine whether this sheet represents the hub site.

    Optional Excel overrides supported:
      - is_hub / hub = TRUE/FALSE
      - role / site_role = HUB
      - hub_site = site number

    Backward-compatible product default: site 1 is the hub.
    """
    site = _site_id(site)

    for table in (md, md_top):
        if table is None or table.empty:
            continue
        row = table.iloc[0]

        explicit = _value(row, ["is_hub", "hub", "HUB"], None)
        if explicit is not None:
            return _is_true(explicit)

        role = _value(row, ["role", "site_role", "site role"], None)
        if role is not None:
            return str(role).strip().lower() == "hub"

        hub_site = _value(row, ["hub_site", "hub site"], None)
        if hub_site is not None:
            return _site_id(hub_site) == site

    return site == "1"


def _effective_vlan_info(row, md, is_hub):
    """
    VLAN 200 is the local ISP transit VLAN in this architecture.
    It is only valid on SW1 at the hub and must never be extended to
    downstream switches or spoke sites.
    """
    vlan_info = _parse_vlan_allocation(row.get("vlan-antall", ""))
    isp_vlan = _get_isp_vlan(md)
    sw_id = _switch_id(row.get("SW"))

    if not (is_hub and sw_id == "1"):
        vlan_info = [(vlan, count) for vlan, count in vlan_info if vlan != isp_vlan]

    return vlan_info


def _site_service_vlans(swi_data, md, is_hub):
    """Return the union of service VLANs used anywhere in this site.

    vlan-antall controls LOCAL access-port allocation only.  A VLAN that is
    present on SW3 must still exist and be allowed across the transit trunks
    on SW2/SW1 so it can reach the site router.

    VLAN 200 is special: _effective_vlan_info() only keeps it on HUB SW1.
    The trunk generator additionally removes it from switch-to-switch trunks.
    """
    vlans = []
    for _, site_row in swi_data.iterrows():
        for vlan, _count in _effective_vlan_info(site_row, md, is_hub):
            vlans.append(vlan)
    return _ordered_unique(vlans)


def _build_port_plan(row, md, md_top, swi_data, site, is_hub=False):
    """Build deterministic physical port allocation for a switch.

    Every switch keeps one dedicated MGMT access port as before.  On HUB SW1
    that port is the TACACS/Syslog server port when both services share an IP.
    If the service IPs differ, HUB SW1 reserves one additional MGMT port for
    Syslog/Security Onion management.  SPAN/RSPAN reserve a local destination
    port on SW1; ERSPAN reserves a mirror destination port only on HUB SW1.
    """
    has_new_num_ports = (
        "num_ports_tot" in row.index
        and not pd.isna(row.get("num_ports_tot"))
        and str(row.get("num_ports_tot")).strip() != ""
    )

    num_ports = _int_value(row, ["num_ports_tot", "num_ports"], 0)
    skipped_ports = _int_value(
        row,
        ["skiped_ports", "skipped_ports"],
        DEFAULT_SKIPPED_PORTS,
    )
    num_downlinks = _int_value(row, ["num_downlink", "num_downlinks"], 0)
    etherchan_ports = _get_etherchannel_ports(md)
    sw_id = _switch_id(row.get("SW"))
    is_primary_switch = sw_id == "1"

    if num_ports <= 0:
        raise ValueError(f"SW{sw_id}: num_ports må være > 0")
    if skipped_ports < 0:
        raise ValueError(f"SW{sw_id}: skiped_ports kan ikke være negativ")
    if num_downlinks < 0:
        raise ValueError(f"SW{sw_id}: num_downlink kan ikke være negativ")
    if has_new_num_ports and skipped_ports >= num_ports:
        raise ValueError(
            f"SW{sw_id}: skiped_ports ({skipped_ports}) må være mindre enn "
            f"num_ports_tot ({num_ports})"
        )

    first_usable = skipped_ports
    if has_new_num_ports:
        last_port = num_ports - 1
        usable_count = num_ports - skipped_ports
        skipped_physical_ports = list(range(0, skipped_ports))
    else:
        last_port = skipped_ports + num_ports - 1
        usable_count = num_ports
        skipped_physical_ports = []

    mgmt_port = first_usable
    tacacs_server = _get_tacacs_server_ip(md_top)
    syslog_server = _get_syslog_server_ip(md_top)
    separate_server_ports = (
        is_hub and is_primary_switch and tacacs_server != syslog_server
    )
    syslog_port = mgmt_port + 1 if separate_server_ports else mgmt_port

    span_mode = _get_span_mode(md)
    if span_mode == "RSPAN" and len(swi_data) < 2:
        raise ValueError(
            f"Site {site}: RSPAN krever minst en downstream-switch. "
            "Sitet har bare SW1; bruk SPAN i stedet."
        )

    local_sensor_port_required = (
        is_primary_switch
        and (
            span_mode in {"SPAN", "RSPAN"}
            or (span_mode == "ERSPAN" and is_hub)
        )
    )

    server_port_count = 1 + (1 if separate_server_ports else 0)
    sensor_port = (
        first_usable + server_port_count if local_sensor_port_required else None
    )
    dedicated_ports = server_port_count + (1 if local_sensor_port_required else 0)

    uplink_member_count = 1 if is_primary_switch else etherchan_ports
    downlink_member_count = num_downlinks * etherchan_ports
    trunk_member_count = uplink_member_count + downlink_member_count
    available_access_slots = usable_count - dedicated_ports - trunk_member_count

    if available_access_slots < 0:
        uplink_desc = "1 router-uplink" if is_primary_switch else f"{etherchan_ports}-ports uplink"
        raise ValueError(
            f"SW{sw_id}: ikke nok porter. {usable_count} brukbare porter, men "
            f"dedikerte MGMT/server/sensor-porter + {uplink_desc} + downlinks krever "
            f"{dedicated_ports + trunk_member_count} porter."
        )

    first_access_port = first_usable + dedicated_ports
    last_access_port = first_access_port + available_access_slots - 1

    uplink_start = last_port - uplink_member_count + 1
    uplink_ports = list(range(uplink_start, last_port + 1))

    downlink_groups = []
    first_downlink_top = uplink_start - 1
    for idx in range(num_downlinks):
        group_end = first_downlink_top - idx * etherchan_ports
        group_start = group_end - etherchan_ports + 1
        downlink_groups.append(list(range(group_start, group_end + 1)))

    vlan_info = _effective_vlan_info(row, md, is_hub)
    allocated_access_ports = sum(max(0, count) for _, count in vlan_info)
    if allocated_access_ports > available_access_slots:
        raise ValueError(
            f"SW{sw_id}: vlan-antall bruker {allocated_access_ports} access-porter, "
            f"men bare {available_access_slots} er tilgjengelige etter "
            f"dedikerte porter/uplink/downlinks."
        )

    expected_free = _value(row, ["num_port_ledig"], None)
    if expected_free is not None:
        try:
            expected_free = int(float(expected_free))
            if expected_free != available_access_slots:
                print(
                    f"ADVARSEL SW{sw_id}: num_port_ledig i Excel er {expected_free}, "
                    f"men generatoren beregner {available_access_slots}."
                )
        except (TypeError, ValueError):
            pass

    return {
        "num_ports": num_ports,
        "skipped_ports": skipped_ports,
        "skipped_physical_ports": skipped_physical_ports,
        "mgmt_port": mgmt_port,
        "tacacs_port": mgmt_port,
        "syslog_port": syslog_port,
        "separate_server_ports": separate_server_ports,
        "span_port": sensor_port,
        "sensor_port": sensor_port,
        "span_mode": span_mode,
        "first_access_port": first_access_port,
        "last_access_port": last_access_port,
        "available_access_slots": available_access_slots,
        "allocated_access_ports": allocated_access_ports,
        "etherchan_ports": etherchan_ports,
        "is_primary_switch": is_primary_switch,
        "uplink_is_etherchannel": (not is_primary_switch and etherchan_ports > 1),
        "downlink_is_etherchannel": etherchan_ports > 1 and num_downlinks > 0,
        "num_downlinks": num_downlinks,
        "uplink_ports": uplink_ports,
        "downlink_groups": downlink_groups,
        "vlan_info": vlan_info,
    }


def create_tacacs_config(md_top):
    my_data = {"config": {}, "network_info": {}}
    tacacs_server = _get_tacacs_server_ip(md_top)
    tacacs_key = md_top.iloc[0].get("tacacs_key", "")

    if pd.isna(tacacs_key) or not str(tacacs_key).strip():
        raise ValueError("TACACS-key mangler i Excel")

    my_data["config"]["aaa new-model"] = []
    my_data["config"]["aaa group server tacacs+ TACACS-GROUP"] = [
        f"server-private {tacacs_server} key {str(tacacs_key).strip()}",
        "ip tacacs source-interface Vlan10",
        "exit",
    ]
    my_data["config"]["aaa authentication login default group TACACS-GROUP local"] = []
    my_data["config"]["aaa authorization exec default group TACACS-GROUP local"] = []
    return my_data


def create_rsyslog_config(md_top):
    my_data = {"config": {}, "network_info": {}}
    rsyslog_server = _get_syslog_server_ip(md_top)

    my_data["config"]["service timestamps log datetime msec show-timezone"] = []
    my_data["config"][f"logging host {rsyslog_server} transport udp port 514"] = []
    my_data["config"]["logging trap informational"] = []
    my_data["config"]["logging source-interface Vlan10"] = []
    return my_data


def enable_ssh(md, domain=SSH_DOMAIN, mgmt_network=None, mgmt_wildcard=None):
    my_data = {"config": {}, "network_info": {}}
    if md.empty:
        return my_data

    row = md.iloc[0]
    username = row.get("brukernavn", "")
    password = row.get("passord", "")
    vty_lines = row.get("vty_lines", "0-4")

    if pd.isna(username) or pd.isna(password):
        return my_data

    username = str(username).strip()
    password = str(password).strip()
    if not username or not password:
        return my_data

    my_data["config"][f"ip domain name {domain}"] = []
    my_data["config"][f"username {username} privilege 15 secret 9 {password}"] = []
    my_data["config"]["crypto key generate rsa general-keys modulus 2048"] = []
    my_data["config"]["ip ssh version 2"] = []

    vty_cfg = ["login authentication default", "transport input ssh"]
    if mgmt_network and mgmt_wildcard:
        my_data["config"]["ip access-list standard SSH-MGMT-ONLY"] = [
            f"permit {mgmt_network} {mgmt_wildcard}",
            "exit",
        ]
        vty_cfg.insert(0, "access-class SSH-MGMT-ONLY in")
    vty_cfg.append("exit")

    my_data["config"][f"line vty {' '.join(x.strip() for x in str(vty_lines).split('-'))}"] = vty_cfg

    return my_data


def _mgmt_access_port_config(description, mgmt_vlan):
    return [
        f"description {description}",
        "ip arp inspection trust",
        "switchport mode access",
        f"switchport access vlan {mgmt_vlan}",
        "switchport port-security",
        "switchport port-security maximum 2",
        "switchport port-security violation restrict",
        "spanning-tree bpduguard enable",
        "spanning-tree portfast",
        "no shutdown",
        "exit",
    ]


def global_config(md, md_top, swi_data, ip_data, vrf_data, is_hub, erspan_context=None):
    info = {"config": {}, "network_info": {}}
    site = _site_id(md.iloc[0]["site"])
    secret = md.iloc[0].get("secret", "")
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )
    erspan_destination = (
        erspan_context.get("destination_ip")
        if span_mode_site == "ERSPAN" and erspan_context
        else None
    )
    erspan_id = (
        int(erspan_context.get("erspan_id", DEFAULT_ERSPAN_ID))
        if erspan_context
        else DEFAULT_ERSPAN_ID
    )

    if span_mode_site == "ERSPAN" and not erspan_destination:
        raise ValueError("ERSPAN-destinasjon kunne ikke utledes fra HUB SW1 MONITORING ip.")

    tacacs_server = _get_tacacs_server_ip(md_top)
    syslog_server = _get_syslog_server_ip(md_top)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        mgmt_ip = row["MGMT ip"]
        gateway = row["gateway"]
        mask = row["mask"]
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg[f"hostname {sw_name}"] = []
        sw_cfg[f"enable secret 9 {secret}"] = []
        sw_cfg.update(create_tacacs_config(md_top)["config"])
        mgmt_network = mgmt_wildcard = None
        try:
            mgmt_net = ipaddress.ip_network(f"{mgmt_ip}/{mask}", strict=False)
            mgmt_network, mgmt_wildcard = str(mgmt_net.network_address), str(mgmt_net.hostmask)
        except (ValueError, TypeError):
            pass
        sw_cfg.update(enable_ssh(md, mgmt_network=mgmt_network, mgmt_wildcard=mgmt_wildcard)["config"])
        sw_cfg.update(create_rsyslog_config(md_top)["config"])

        sw_cfg[f"vlan {mgmt_vlan}"] = [f"name MGMT_VLAN_{mgmt_vlan}", "exit"]
        sw_cfg[f"interface vlan {mgmt_vlan}"] = [
            f"ip address {mgmt_ip} {mask}",
            "no shutdown",
            "exit",
        ]

        monitoring_ip = None
        if span_mode_site == "ERSPAN":
            monitoring_ip = _get_switch_monitoring_ip(row, monitoring, site, sw_id)
            monitoring_vlan = monitoring["vlan"]
            monitoring_mask = monitoring["mask"]
            monitoring_gateway = monitoring["gateway"]

            sw_cfg["ip routing"] = []
            sw_cfg[f"vlan {monitoring_vlan}"] = ["name MONITORING", "exit"]
            sw_cfg[f"interface vlan {monitoring_vlan}"] = [
                f"ip address {monitoring_ip} {monitoring_mask}",
                "no shutdown",
                "exit",
            ]
            sw_cfg[f"ip route 0.0.0.0 0.0.0.0 {gateway}"] = []

            destination_ip = ipaddress.ip_address(erspan_destination)
            if destination_ip not in monitoring["network"]:
                sw_cfg[
                    f"ip route {erspan_destination} 255.255.255.255 {monitoring_gateway}"
                ] = []

        # Preserve the existing dedicated MGMT port on every switch.  On HUB SW1
        # it becomes the physical server handoff for TACACS/Syslog.
        if is_hub and plan["is_primary_switch"]:
            if tacacs_server == syslog_server:
                mgmt_desc = "Dedicated TACACS/Syslog management server port"
            else:
                mgmt_desc = "Dedicated TACACS management server port"
        else:
            mgmt_desc = f"Dedicated management access port for VLAN {mgmt_vlan}"

        sw_cfg[f"interface {intf_prefix}{plan['tacacs_port']}"] = _mgmt_access_port_config(
            mgmt_desc, mgmt_vlan
        )

        if is_hub and plan["is_primary_switch"] and plan["separate_server_ports"]:
            sw_cfg[f"interface {intf_prefix}{plan['syslog_port']}"] = _mgmt_access_port_config(
                "Dedicated Syslog/Security Onion management server port", mgmt_vlan
            )

        # ERSPAN destination port is a true mirror destination.  It is not an
        # access port and carries no IP/VLAN configuration toward the sensor NIC.
        if (
            span_mode_site == "ERSPAN"
            and is_hub
            and plan["is_primary_switch"]
            and plan["sensor_port"] is not None
        ):
            sw_cfg[f"interface {intf_prefix}{plan['sensor_port']}"] = [
                "description Dedicated Security Onion ERSPAN mirror destination port",
                "no shutdown",
                "exit",
            ]

        span_mode = plan["span_mode"]
        span_vlans = _ordered_unique([mgmt_vlan, *[v for v, _ in plan["vlan_info"]]])
        if span_mode == "ERSPAN" and monitoring is not None:
            span_vlans = [v for v in span_vlans if v != monitoring["vlan"]]

        if span_mode == "SPAN" and plan["span_port"] is not None:
            sw_cfg[f"interface {intf_prefix}{plan['span_port']}"] = [
                "description Dedicated local SPAN destination port for IDS/IPS",
                "no shutdown",
                "exit",
            ]
            sw_cfg[f"monitor session 1 source vlan {','.join(map(str, span_vlans))} both"] = []
            sw_cfg[f"monitor session 1 destination interface {intf_prefix}{plan['span_port']}"] = []

        elif span_mode == "RSPAN":
            rspan_vlan = _get_rspan_vlan(md)
            if plan["is_primary_switch"]:
                if plan["span_port"] is None:
                    raise ValueError(
                        f"SW{sw_id}: RSPAN krever en lokal sensorport på primærswitchen."
                    )
                sw_cfg[f"interface {intf_prefix}{plan['span_port']}"] = [
                    "description Dedicated RSPAN destination port for IDS/IPS",
                    "no shutdown",
                    "exit",
                ]
                sw_cfg[f"monitor session 1 source remote vlan {rspan_vlan}"] = []
                sw_cfg[f"monitor session 1 destination interface {intf_prefix}{plan['span_port']}"] = []
            else:
                sw_cfg[f"monitor session 1 source vlan {','.join(map(str, span_vlans))} both"] = []
                sw_cfg[f"monitor session 1 destination remote vlan {rspan_vlan}"] = []

        elif span_mode == "ERSPAN":
            if is_hub and plan["is_primary_switch"]:
                # One common ERSPAN-ID lets a single destination session receive
                # mirrored traffic from all remote ERSPAN source switches.
                sw_cfg["monitor session 1 type erspan-destination"] = [
                    "description ERSPAN-TO-SECURITY-ONION",
                    f"destination interface {intf_prefix}{plan['sensor_port']}",
                    "source",
                    f"erspan-id {erspan_id}",
                    f"ip address {erspan_destination}",
                    "no shutdown",
                    "exit",
                    "exit",
                ]
                print(
                    f"ADVARSEL Site {site} SW1: ERSPAN destination-porten kan bare tilhøre én "
                    "monitor-session. Lokal-only trafikk på HUB SW1 speiles derfor ikke i ERSPAN-modus. "
                    "Remote sites speiler klient-ingress og router-uplink-ingress for å redusere duplikater."
                )
            else:
                # Avoid duplicate ERSPAN copies across a multi-switch site:
                # - mirror ingress on real client access ports
                # - on SW1, also mirror ingress from the site router/WAN
                # - never mirror switch-to-switch trunks
                source_lines = _erspan_source_interface_lines(plan, intf_prefix)

                if source_lines:
                    sw_cfg["monitor session 1 type erspan-source"] = [
                        f"description ERSPAN-SITE-{site}-SW{sw_id}",
                        *source_lines,
                        "destination",
                        f"ip address {erspan_destination}",
                        f"erspan-id {erspan_id}",
                        f"origin ip-address {monitoring_ip}",
                        "ip ttl 32",
                        "exit",
                        "no shutdown",
                        "exit",
                    ]
                else:
                    print(
                        f"INFO Site {site} SW{sw_id}: ingen ERSPAN-source opprettet; "
                        "switchen har ingen klient-accessporter å speile."
                    )

        if span_mode_site != "ERSPAN":
            sw_cfg[f"ip default-gateway {gateway}"] = []
        sw_cfg[f"ntp server {gateway}"] = []

    return info


def config_vlan(swi_data, site, md, md_top, ip_data, vrf_data, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )

    # IMPORTANT:
    # vlan-antall is a LOCAL access-port requirement, not a statement that the
    # VLAN exists only on that switch.  Build a site-wide union so intermediate
    # switches can actually forward VLANs used farther downstream.
    site_service_vlans = _site_service_vlans(swi_data, md, is_hub)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)
        local_vlan_info = plan["vlan_info"]

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        sw_cfg["vlan 999"] = ["name NATIVE_UBRUKT", "exit"]

        # Create every site service VLAN on every switch that may need to
        # transport it.  ISP VLAN 200 remains local to HUB SW1 only.
        vlans_to_create = [
            vlan
            for vlan in site_service_vlans
            if vlan != isp_vlan or (is_hub and plan["is_primary_switch"])
        ]
        for vlan in vlans_to_create:
            sw_cfg[f"vlan {vlan}"] = [f"name VLAN_{vlan}", "exit"]

        if span_mode_site == "ERSPAN":
            mon_vlan = monitoring["vlan"]
            service_vlans = {mgmt_vlan, *site_service_vlans}
            if mon_vlan in service_vlans:
                raise ValueError(
                    f"SW{sw_id}: MONITORING VLAN {mon_vlan} kolliderer med et eksisterende tjeneste-VLAN."
                )
            sw_cfg[f"vlan {mon_vlan}"] = ["name MONITORING", "exit"]

        if plan["span_mode"] == "RSPAN":
            rspan_vlan = _get_rspan_vlan(md)
            service_vlans = {mgmt_vlan, *site_service_vlans}
            if rspan_vlan in service_vlans:
                raise ValueError(
                    f"SW{sw_id}: RSPAN-VLAN {rspan_vlan} kolliderer med et tjeneste-VLAN."
                )
            sw_cfg[f"vlan {rspan_vlan}"] = ["name RSPAN_MONITOR", "remote-span", "exit"]

        # LOCAL access-port allocation still comes only from this switch's
        # vlan-antall entry.
        current_port = plan["first_access_port"]
        for vlan, count in local_vlan_info:
            if count <= 0:
                continue

            ports = list(range(current_port, current_port + count))
            current_port += count
            key = _interface_key(intf_prefix, ports)

            port_cfg = [
                f"description access port for VLAN {vlan}",
            ]
            if vlan == isp_vlan:
                # ISP transit port: infrastructure hand-off, not a client access port.
                port_cfg.extend([
                    "switchport mode access",
                    f"switchport access vlan {vlan}",
                    "no shutdown",
                    "exit",
                ])
                sw_cfg[key] = port_cfg
                continue

            port_cfg.extend([
                "switchport mode access",
                f"switchport access vlan {vlan}",
                "switchport port-security",
                "switchport port-security maximum 2",
                "switchport port-security violation restrict",
                "ip verify source",
                "spanning-tree bpduguard enable",
                "spanning-tree portfast",
                "no shutdown",
                "exit",
            ])
            sw_cfg[key] = port_cfg

    return info


def _trunk_config(vlans, description, channel_group=None, trusted=True):
    lines = [
        description,
        "switchport trunk encapsulation dot1q",
        "switchport trunk native vlan 999",
        "switchport mode trunk",
        f"switchport trunk allowed vlan {','.join(map(str, vlans))}",
    ]
    if trusted:
        # Trust only flows toward the DHCP server / infrastructure side (uplinks).
        lines.extend(["ip dhcp snooping trust", "ip arp inspection trust"])
    if channel_group is not None:
        lines.append(f"channel-group {channel_group} mode active")
    lines.extend(["no shutdown", "exit"])
    return lines


def config_trunk_and_dchp_snooping(swi_data, site, md, md_top, ip_data, vrf_data, is_hub):
    info = {"config": {}, "network_info": {}}
    isp_vlan = _get_isp_vlan(md)
    span_mode_site = _get_span_mode(md)
    monitoring = (
        _get_monitoring_service(ip_data, vrf_data, site)
        if span_mode_site == "ERSPAN"
        else None
    )

    # VLANs required anywhere in the site must traverse the intermediate
    # switch trunks, even when a particular switch has zero local access ports
    # in that VLAN.
    site_service_vlans = _site_service_vlans(swi_data, md, is_hub)

    for _, row in swi_data.iterrows():
        sw_id = _switch_id(row["SW"])
        mgmt_vlan = _int_value(row, ["MGMT Vlan"], 10)
        intf_prefix = str(row["intf_prefix"]).strip()
        plan = _build_port_plan(row, md, md_top, swi_data, site, is_hub)

        sw_name = f"SW{sw_id}-SITE-{site}"
        info["config"].setdefault(sw_name, {})
        sw_cfg = info["config"][sw_name]

        # Do NOT derive trunk VLANs from this switch's vlan-antall.
        # vlan-antall only says how many LOCAL access ports the switch needs.
        all_trunk_vlans = _ordered_unique([mgmt_vlan, *site_service_vlans, 999])
        if span_mode_site == "ERSPAN":
            all_trunk_vlans = _ordered_unique([*all_trunk_vlans[:-1], monitoring["vlan"], 999])

        # VLAN 200 is local ISP transit in this architecture and must not be extended downstream.
        downlink_vlans = [v for v in all_trunk_vlans if v != isp_vlan]

        # RSPAN is a site-local Layer-2 transport VLAN. It is carried only on
        # switch-to-switch trunks, never on the SW1 -> router trunk.
        switch_link_vlans = list(downlink_vlans)
        if plan["span_mode"] == "RSPAN":
            print("ADVARSEL:")
            print("RSPAN samler trafikk fra downstream-switcher.")
            print("Lokal trafikk på RSPAN-destination-switchen SW1 speiles ikke.")
            rspan_vlan = _get_rspan_vlan(md)
            if rspan_vlan in all_trunk_vlans:
                raise ValueError(
                    f"SW{sw_id}: RSPAN-VLAN {rspan_vlan} kolliderer med et eksisterende VLAN."
                )
            switch_link_vlans = _ordered_unique([*switch_link_vlans, rspan_vlan])

        # No snooping/DAI on ISP transit, RSPAN or blackhole/native VLAN 999.
        excluded_inspection_vlans = {isp_vlan, 999}
        if span_mode_site == "ERSPAN":
            excluded_inspection_vlans.add(monitoring["vlan"])
        inspection_vlans = [v for v in all_trunk_vlans if v not in excluded_inspection_vlans]
        if inspection_vlans:
            sw_cfg["ip dhcp snooping"] = []
            sw_cfg[f"ip dhcp snooping vlan {','.join(map(str, inspection_vlans))}"] = []
            sw_cfg["no ip dhcp snooping information option"] = []
            sw_cfg[f"ip arp inspection vlan {','.join(map(str, inspection_vlans))}"] = []

        # UPLINK:
        # - SW1: exactly one physical trunk to the site router.
        # - SW2+: EtherChannel-sized uplink toward the upstream switch when
        #   etherchan_num > 1; otherwise a normal physical trunk.
        if plan["uplink_is_etherchannel"]:
            uplink_po = 1
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                switch_link_vlans,
                f"description UPLINK EtherChannel member(s) - Port-channel{uplink_po}",
                channel_group=uplink_po,
                trusted=True,
            )
            sw_cfg[f"interface Port-channel{uplink_po}"] = _trunk_config(
                switch_link_vlans,
                f"description UPLINK Port-channel{uplink_po} toward upstream switch "
                f"for VLAN {','.join(map(str, switch_link_vlans))}",
                trusted=True,
            )
            first_downlink_channel = 2
        else:
            uplink_vlans = all_trunk_vlans if plan["is_primary_switch"] else switch_link_vlans
            uplink_desc = (
                "UPLINK trunk to site router"
                if plan["is_primary_switch"]
                else "UPLINK trunk toward upstream switch"
            )
            sw_cfg[_interface_key(intf_prefix, plan["uplink_ports"])] = _trunk_config(
                uplink_vlans,
                f"description {uplink_desc} for VLAN {','.join(map(str, uplink_vlans))}",
                trusted=True,
            )
            first_downlink_channel = 1

        # DOWNLINKS:
        # Channel-group numbers are local. SW1 starts with Po1. On downstream
        # switches Po1 is reserved for the uplink, so their downlinks start at Po2.
        if plan["downlink_is_etherchannel"]:
            for offset, ports in enumerate(plan["downlink_groups"]):
                channel_id = first_downlink_channel + offset
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK EtherChannel member(s) - Port-channel{channel_id}",
                    channel_group=channel_id,
                    trusted=False,
                )
                sw_cfg[f"interface Port-channel{channel_id}"] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK Port-channel{channel_id} for VLAN "
                    f"{','.join(map(str, switch_link_vlans))}",
                    trusted=False,
                )
        else:
            for idx, ports in enumerate(plan["downlink_groups"], start=1):
                sw_cfg[_interface_key(intf_prefix, ports)] = _trunk_config(
                    switch_link_vlans,
                    f"description DOWNLINK trunk {idx} for VLAN "
                    f"{','.join(map(str, switch_link_vlans))}",
                    trusted=False,
                )

        # New num_ports_tot model: skipped ports are real physical interfaces,
        # so explicitly blackhole/shut them instead of leaving them in VLAN 1.
        if plan["skipped_physical_ports"]:
            sw_cfg[_interface_key(intf_prefix, plan["skipped_physical_ports"])] = [
                "description UBRUKT - SKIPPED/RESERVED",
                "switchport mode access",
                "switchport access vlan 999",
                "shutdown",
                "exit",
            ]

        # Remaining unassigned access slots are blackholed and shut down.
        unused_start = plan["first_access_port"] + plan["allocated_access_ports"]
        unused_end = plan["last_access_port"]
        if unused_start <= unused_end:
            unused_ports = list(range(unused_start, unused_end + 1))
            sw_cfg[_interface_key(intf_prefix, unused_ports)] = [
                "description UBRUKT - BLACKHOLE VLAN",
                "switchport mode access",
                "switchport access vlan 999",
                "shutdown",
                "exit",
            ]

    return info


def update_site_config(data, swi_data, sn, conf):
    for raw_sw_id in swi_data["SW"].dropna().unique():
        sw_id = _switch_id(raw_sw_id)
        sw_name = f"SW{sw_id}-SITE-{sn}"
        data[f"site {sn}"]["config"].setdefault(sw_name, {})
        data[f"site {sn}"]["config"][sw_name].update(conf["config"][sw_name])
    return data


def fetch_site_data(config_file):
    try:
        with open(config_file, "r") as f:
            try:
                return json.load(f)
            except json.JSONDecodeError:
                return {}
    except FileNotFoundError:
        return {}


def _refresh_ssh_acls(data, swi_data, sn):
    """
    Record this site's MGMT subnet and re-sync the SSH-MGMT-ONLY ACL on every
    switch across every site processed so far, so newly added sites are
    automatically permitted everywhere (not just locally).
    """
    mgmt_networks = data.setdefault("_mgmt_networks", {})

    row = swi_data.iloc[0]
    try:
        net = ipaddress.ip_network(f"{row.get('MGMT ip')}/{row.get('mask')}", strict=False)
        mgmt_networks[sn] = [str(net.network_address), str(net.hostmask)]
    except (ValueError, TypeError):
        pass

    if not mgmt_networks:
        return

    acl_lines = [f"permit {network} {wildcard}" for network, wildcard in sorted(mgmt_networks.values())]
    acl_lines.append("exit")

    for site_key, site_val in data.items():
        if site_key == "_mgmt_networks":
            continue
        for sw_cfg in site_val.get("config", {}).values():
            if "ip access-list standard SSH-MGMT-ONLY" in sw_cfg:
                sw_cfg["ip access-list standard SSH-MGMT-ONLY"] = acl_lines


def create_site_sw_config(file, sheet, config_file, erspan_context=None):
    sheet_data = read_sheet(file, sheet)
    data = fetch_site_data(config_file)

    md = sheet_data["md"]
    md_top = sheet_data["md_top"]
    swi_data = sheet_data["swi_data"]
    ip_data = sheet_data.get("ip_data")
    vrf_data = sheet_data.get("vrf_data")

    sn = _site_id(md.iloc[0]["site"])
    is_hub = _is_hub_site(md, md_top, sn)
    data[f"site {sn}"] = {"config": {}}

    vlan_conf = config_vlan(swi_data, sn, md, md_top, ip_data, vrf_data, is_hub)
    data = update_site_config(data, swi_data, sn, vlan_conf)

    global_conf = global_config(md, md_top, swi_data, ip_data, vrf_data, is_hub, erspan_context)
    data = update_site_config(data, swi_data, sn, global_conf)

    trunk_conf = config_trunk_and_dchp_snooping(swi_data, sn, md, md_top, ip_data, vrf_data, is_hub)
    data = update_site_config(data, swi_data, sn, trunk_conf)

    _refresh_ssh_acls(data, swi_data, sn)

    with open(config_file, "w") as f:
        json.dump(data, f, indent=4)

    return data


def config_to_text(data, indent=0):
    lines = []
    prefix = "    " * indent

    if isinstance(data, dict):
        for key, value in data.items():
            lines.append("!")
            lines.append(prefix + key)
            lines.extend(config_to_text(value, indent + 1))
    elif isinstance(data, list):
        for value in data:
            if isinstance(value, str):
                lines.append(prefix + value)
            else:
                lines.extend(config_to_text(value, indent))
    elif isinstance(data, str):
        lines.append(prefix + data)

    return lines


def create_or_update_config_files(data):
    os.makedirs("siteSwichTextConfigs", exist_ok=True)

    for site, site_data in data.items():
        if site == "_mgmt_networks":
            continue
        site_dir = f"siteSwichTextConfigs/{site}"
        os.makedirs(site_dir, exist_ok=True)

        for sw_name, config in site_data["config"].items():
            text = config_to_text(config)
            with open(f"{site_dir}/{sw_name}.txt", "w", encoding="utf-8") as f:
                f.write("\n".join(text))

    print("Text versjon av switch-configene er lagret i siteSwichTextConfigs/")


def create_sw_configs_main(file, config_file="site_switch_config.json"):
    sites_sheets = load_workbook(file, read_only=True).sheetnames

    # Fail fast before any config files are generated if enabled sites mix
    # SPAN/RSPAN/ERSPAN modes. Blank span_mode is allowed and means OFF.
    _validate_span_modes_for_workbook(file, sites_sheets)
    _validate_management_servers_for_workbook(file, sites_sheets)
    erspan_context = _validate_erspan_architecture_for_workbook(file, sites_sheets)

    data = {}
    for sheet in sites_sheets:
        data = create_site_sw_config(file, sheet, config_file, erspan_context)

    create_or_update_config_files(data)


def main():
    file = sys.argv[1]
    config_file = "site_switch_config.json" if len(sys.argv) < 3 else sys.argv[2]
    create_sw_configs_main(file, config_file)


if __name__ == "__main__":
    main()